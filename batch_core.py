# -*- coding: utf-8 -*-
"""
Ver.4.3 一括分類の処理本体。自動録音機のフォルダを丸ごと処理し、Excel一覧表を作る。

  ・Perch（Ver.4.0系）とBirdNET（Ver.4.1系）の両方で判定する
  ・どちらも「コウモリ22種＋ノイズ（非コウモリ音）」の23クラス分類器（Ver.4.3で再学習）
  ・1ファイル処理するごとに途中経過CSVへ追記し、中断しても続きから再開できる

コマンドラインから使う場合（venv_perch環境）:
  python batch_core.py <対象フォルダ> [出力Excelパス]
画面から使う場合は batch_app.py（start_batch_app.bat）を参照。
"""
import csv
import datetime as dt
import io
import os
import pathlib
import re
import shutil
import sys
import tempfile
import time
import warnings
import zipfile
warnings.filterwarnings("ignore")

import joblib
import numpy as np
import pandas as pd
import soundfile as sf
from scipy import signal as scipy_signal

BASE_DIR  = pathlib.Path(__file__).parent
# ローカル（Ver4.3_batch/models/）とGitHubリポジトリ直下（ver43_batch/、Web版と共用）の両方に対応
MODEL_DIR = next((d for d in (BASE_DIR / "models", BASE_DIR / "ver43_batch")
                  if (d / "perch_classifier_noise.joblib").exists()), BASE_DIR / "models")
TARGET_SR = 44100        # ピッチシフト後（サンプルレートのラベル付け替え、Ver.4.0/4.1と同じ）
NOISE = "ノイズ"
TOP_K = 3
BATCH_FILES = 16         # 埋め込みモデルにまとめて渡すファイル数

LOW_CONF = 0.4
LOW_SR_HZ = 192000       # これ未満はフルスペクトラム超音波録音でない可能性
LONG_SEC = 30.0          # これより長いと平均化で短い通過音が埋もれる可能性

SINGLE_LOCATION_SPECIES = {"オヒキコウモリ", "キタクビワコウモリ", "コヤマコウモリ",
                           "ドーベントンコウモリ", "ヒメヒナコウモリ"}

DATETIME_RE = re.compile(r"(20\d{2})(\d{2})(\d{2})[_\-T]?(\d{2})(\d{2})(\d{2})")

COLUMNS = ["No", "ファイル名", "フォルダ", "録音日時", "日時の取得元", "録音長(秒)",
           "サンプルレート(kHz)", "パルス検出比率", "総合判定", "要確認", "確認理由",
           "Perch第1候補", "Perch確信度", "Perch第2候補", "Perch第2確信度", "Perch第3候補", "Perch第3確信度",
           "BirdNET第1候補", "BirdNET確信度", "BirdNET第2候補", "BirdNET第2確信度",
           "BirdNET第3候補", "BirdNET第3確信度", "エラー", "フルパス"]


# ─── モデル ────────────────────────────────────────────
def load_models():
    import bioacoustics_model_zoo as bmz
    return {
        "Perch":   (bmz.Perch2ONNX(headless=True), joblib.load(MODEL_DIR / "perch_classifier_noise.joblib")),
        "BirdNET": (bmz.BirdNET(), joblib.load(MODEL_DIR / "birdnet_classifier_noise.joblib")),
    }


# ─── ファイル情報 ───────────────────────────────────────
def list_wavs(folder, recursive=True):
    folder = pathlib.Path(folder)
    it = folder.rglob("*") if recursive else folder.glob("*")
    return sorted(p for p in it if p.is_file() and p.suffix.lower() == ".wav")


def recording_datetime(path, trust_mtime=True):
    """N_20200915_063713.wav 形式ならファイル名から、それ以外はファイル更新日時で代用。
    Web版で個別アップロードされたファイルは更新日時が失われるので trust_mtime=False で「不明」にする"""
    m = DATETIME_RE.search(path.stem)
    if m:
        try:
            return dt.datetime(*map(int, m.groups())), "ファイル名"
        except ValueError:
            pass
    if not trust_mtime:
        return None, "不明"
    return dt.datetime.fromtimestamp(path.stat().st_mtime), "更新日時で代用"


def pulse_ratio(data, sr, n_fft=2048, freq_min_hz=15000, threshold_db=15.0):
    """Ver.4.0のQC（rescan_no_pulses_fixed.py）と同じ指標。15kHz以上で、フレーム最大値が
    ノイズフロア(10パーセンタイル)より15dB以上高いフレームの割合"""
    if sr / 2 <= freq_min_hz or len(data) < n_fft:
        return float("nan")
    hop = n_fft // 4
    f, t, S = scipy_signal.spectrogram(data, fs=sr, nperseg=n_fft, noverlap=n_fft - hop,
                                       window="hann", scaling="spectrum")
    fm = (10.0 * np.log10(S + 1e-10))[f >= freq_min_hz].max(axis=0)
    return float(np.mean(fm > np.percentile(fm, 10) + threshold_db))


# ─── 判定 ─────────────────────────────────────────────
def top_k(clf, feature):
    probs = clf.predict_proba(feature.reshape(1, -1))[0]
    order = np.argsort(probs)[::-1][:TOP_K]
    return [(clf.classes_[i], float(probs[i])) for i in order]


def summarize(perch, birdnet):
    """2方式の第1候補から総合判定と要確認の理由を決める"""
    (p_sp, p_conf), (b_sp, b_conf) = perch[0], birdnet[0]
    reasons = []
    if p_sp == b_sp:
        final = p_sp
    elif NOISE in (p_sp, b_sp):
        final = "要確認"
        reasons.append("コウモリ/ノイズで2方式不一致")
    else:
        final = "要確認"
        reasons.append("種名が2方式で不一致")
    for name, sp, conf in (("Perch", p_sp, p_conf), ("BirdNET", b_sp, b_conf)):
        if conf < LOW_CONF:
            reasons.append(f"{name}確信度が低い({conf:.0%})")
    for sp in {p_sp, b_sp} & SINGLE_LOCATION_SPECIES:
        reasons.append(f"{sp}は単一調査地データのみで学習")
    return final, reasons


# ─── Web版のアップロード展開 ─────────────────────────────
def save_uploads(uploads, workdir):
    """アップロード（.name と .getvalue() を持つオブジェクト）を作業フォルダに保存する。
    ZIPは展開し、ZIP内に記録された日時を更新日時として復元する（録音日時の代用に使うため）"""
    workdir = pathlib.Path(workdir)
    for up in uploads:
        if up.name.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(up.getvalue())) as zf:
                for info in zf.infolist():
                    if info.is_dir() or not info.filename.lower().endswith(".wav"):
                        continue
                    # Windows標準の圧縮は日本語ファイル名をcp932で記録する（UTF-8フラグ無し）
                    name = info.filename
                    if not info.flag_bits & 0x800:
                        try:
                            name = info.filename.encode("cp437").decode("cp932")
                        except (UnicodeEncodeError, UnicodeDecodeError):
                            pass
                    rel = pathlib.PurePosixPath(name)
                    if rel.is_absolute() or ".." in rel.parts:
                        continue
                    out = workdir / rel
                    out.parent.mkdir(parents=True, exist_ok=True)
                    out.write_bytes(zf.read(info))
                    ts = time.mktime(dt.datetime(*info.date_time).timetuple())
                    os.utime(out, (ts, ts))
        else:
            loose = workdir / "_個別アップロード"
            loose.mkdir(exist_ok=True)
            (loose / pathlib.Path(up.name).name).write_bytes(up.getvalue())


# ─── 本体 ─────────────────────────────────────────────
def default_output_path(folder):
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    return pathlib.Path(folder) / f"コウモリ一括判定_{stamp}.xlsx"


def progress_path_for(xlsx_path):
    p = pathlib.Path(xlsx_path)
    return p.with_name(p.stem + "_途中経過.csv")


def load_progress(progress_csv):
    if not pathlib.Path(progress_csv).exists():
        return []
    with open(progress_csv, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _append_rows(progress_csv, rows):
    new = not pathlib.Path(progress_csv).exists()
    with open(progress_csv, "a", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        if new:
            w.writeheader()
        w.writerows(rows)


def _embed_mean(model, paths):
    """ファイルごとに窓の埋め込みを平均（学習時と同じ平均プーリング）"""
    emb = model.embed(paths, progress_bar=False)
    return emb.groupby(level="file").mean()


def process_batch(models, paths, folder, start_no, tmpdir, trust_mtime=True):
    rows, shifted = [], {}
    for k, p in enumerate(paths):
        row = {c: "" for c in COLUMNS}
        row.update({"No": start_no + k, "ファイル名": p.name, "フルパス": str(p),
                    "フォルダ": str(p.parent.relative_to(folder)) if p.parent != folder else "."})
        try:
            when, src = recording_datetime(p, trust_mtime and "_個別アップロード" not in p.parts)
            row.update({"録音日時": when.strftime("%Y-%m-%d %H:%M:%S") if when else "", "日時の取得元": src})
            data, sr = sf.read(str(p), dtype="float32")
            if data.ndim > 1:
                data = data[:, 0]
            row.update({"録音長(秒)": round(len(data) / sr, 3), "サンプルレート(kHz)": round(sr / 1000, 1)})
            pr = pulse_ratio(data, sr)
            row["パルス検出比率"] = "" if np.isnan(pr) else round(pr, 4)
            tmp = str(tmpdir / f"{start_no + k}.wav")
            sf.write(tmp, data, TARGET_SR)
            shifted[tmp] = (row, sr, len(data) / sr)
        except Exception as e:
            row.update({"総合判定": "読込エラー", "要確認": "要確認", "エラー": str(e)})
        rows.append(row)

    if shifted:
        tmp_paths = list(shifted)
        results = {}
        for name, (model, clf) in models.items():
            try:
                emb = _embed_mean(model, tmp_paths)
            except Exception:
                # 1ファイルの不具合でバッチ全体が落ちないよう、1件ずつやり直す
                parts = []
                for tp in tmp_paths:
                    try:
                        parts.append(_embed_mean(model, [tp]))
                    except Exception as e:
                        shifted[tp][0]["エラー"] += f"{name}: {e} "
                emb = pd.concat(parts) if parts else pd.DataFrame()
            for tp in emb.index:
                results.setdefault(tp, {})[name] = top_k(clf, emb.loc[tp].to_numpy(dtype=np.float64))

        for tp, (row, sr, dur) in shifted.items():
            r = results.get(tp, {})
            if len(r) < 2:
                row.update({"総合判定": "判定エラー", "要確認": "要確認"})
                continue
            for name in ("Perch", "BirdNET"):
                for i, (sp, conf) in enumerate(r[name]):
                    sfx = "" if i == 0 else f"第{i + 1}"
                    row[f"{name}第{i + 1}候補"] = sp
                    row[f"{name}{sfx}確信度"] = round(conf, 4)
            final, reasons = summarize(r["Perch"], r["BirdNET"])
            if sr < LOW_SR_HZ:
                reasons.append(f"サンプルレートが低い({sr / 1000:.0f}kHz)")
            if dur > LONG_SEC:
                reasons.append(f"長時間録音({dur:.0f}秒)：平均化で短い通過音が埋もれる可能性")
            row.update({"総合判定": final, "要確認": "要確認" if reasons else "",
                        "確認理由": "／".join(reasons)})
    return rows


def run(folder, xlsx_path=None, recursive=True, models=None, on_progress=None, trust_mtime=True):
    """フォルダを処理してExcelを書き出す。途中経過CSVがあれば続きから再開する。
    on_progress(done, total, elapsed_sec) を渡すと進捗を通知する。"""
    folder = pathlib.Path(folder)
    xlsx_path = pathlib.Path(xlsx_path) if xlsx_path else default_output_path(folder)
    progress_csv = progress_path_for(xlsx_path)

    files = list_wavs(folder, recursive)
    done_rows = load_progress(progress_csv)
    done = {r["フルパス"] for r in done_rows}
    todo = [p for p in files if str(p) not in done]
    total = len(files)
    if models is None:
        models = load_models()

    t0 = time.time()
    tmpdir = pathlib.Path(tempfile.mkdtemp(prefix="batid_"))
    try:
        for i in range(0, len(todo), BATCH_FILES):
            rows = process_batch(models, todo[i:i + BATCH_FILES], folder, len(done) + i + 1, tmpdir, trust_mtime)
            _append_rows(progress_csv, rows)
            for f in tmpdir.glob("*.wav"):
                f.unlink()
            if on_progress:
                on_progress(len(done) + i + len(rows), total, time.time() - t0)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

    df = pd.DataFrame(load_progress(progress_csv), columns=COLUMNS)
    write_excel(df, xlsx_path, folder)
    return xlsx_path, df


# ─── Excel出力 ─────────────────────────────────────────
NUMERIC_COLS = ["No", "録音長(秒)", "サンプルレート(kHz)", "パルス検出比率", "Perch確信度", "Perch第2確信度",
                "Perch第3確信度", "BirdNET確信度", "BirdNET第2確信度", "BirdNET第3確信度"]

EXPLANATION = [
    ("このファイルについて", "日本産コウモリ音声識別 Ver.4.3（一括分類版）による自動判定の一覧表です。研究用の試作であり、最終判断は専門家が行ってください。"),
    ("判定方式", "Perch（Google、Ver.4.0系）とBirdNET（Ver.4.1系）の音響埋め込みに線形分類器を載せた2方式。どちらもコウモリ22種＋「ノイズ（非コウモリ音）」の23クラスから選びます。"),
    ("総合判定", "2方式の第1候補が一致すればその名前、一致しなければ「要確認」。"),
    ("要確認・確認理由", "2方式の不一致、確信度40%未満、単一調査地データのみの種、低サンプルレート、長時間録音のいずれかに該当すると付きます。"),
    ("パルス検出比率", "15kHz以上で背景より15dB以上強いフレームの割合（判定とは独立の参考値）。0に近いのにコウモリと判定された行は要注意。"),
    ("録音日時", "ファイル名に「20200915_063713」のような日時があればそこから、無ければファイル更新日時で代用（日時の取得元列を参照）。"),
    ("精度の目安（Ver.4.0/4.1、未知の調査地）", "複数地点データがある17種で Perch 49.6% / BirdNET 45.8%。種名は「候補」として扱ってください。"),
    ("ノイズクラスの限界", "ノイズの学習データは虫・雑音・川の音49録音（6調査地）と既存録音の無音区間のみ。雨・風・鳥・人工音などは十分に学習していません。"),
    ("単一調査地の5種", "オヒキ・キタクビワ・コヤマ・ドーベントン・ヒメヒナは1地点のデータのみで学習しており、他地点での性能は未検証です。"),
]


def write_excel(df, xlsx_path, folder):
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter

    df = df.copy()
    for c in NUMERIC_COLS:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["No"] = df["No"].astype("Int64")
    df = df.sort_values("No")

    summary = (df["総合判定"].value_counts().rename_axis("総合判定").reset_index(name="ファイル数"))
    summary["うち要確認"] = summary["総合判定"].map(df[df["要確認"] == "要確認"]["総合判定"].value_counts()).fillna(0).astype(int)
    df["録音日"] = df["録音日時"].str[:10]
    by_date = pd.crosstab(df["録音日"], df["総合判定"]).reset_index()
    df = df.drop(columns="録音日")

    xlsx_path = pathlib.Path(xlsx_path)
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="判定一覧", index=False)
        summary.to_excel(xw, sheet_name="判定別集計", index=False)
        by_date.to_excel(xw, sheet_name="録音日別集計", index=False)
        info = pd.DataFrame(EXPLANATION, columns=["項目", "説明"])
        info.loc[len(info)] = ["対象フォルダ", str(folder)]
        info.loc[len(info)] = ["作成日時", dt.datetime.now().strftime("%Y-%m-%d %H:%M")]
        info.to_excel(xw, sheet_name="説明", index=False)

        yellow = PatternFill("solid", fgColor="FFF2CC")
        grey = PatternFill("solid", fgColor="E7E6E6")
        bold = Font(bold=True)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = bold
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(8, width * 1.6), 60)

        ws = xw.book["判定一覧"]
        idx = {c: i for i, c in enumerate(COLUMNS)}
        for row in ws.iter_rows(min_row=2):
            if row[idx["要確認"]].value == "要確認":
                fill = yellow
            elif row[idx["総合判定"]].value == NOISE:
                fill = grey
            else:
                continue
            for cell in row:
                cell.fill = fill
        for col in ("Perch確信度", "Perch第2確信度", "Perch第3確信度", "BirdNET確信度",
                    "BirdNET第2確信度", "BirdNET第3確信度", "パルス検出比率"):
            for cell in ws[get_column_letter(idx[col] + 1)][1:]:
                cell.number_format = "0.0%"
        ws_info = xw.book["説明"]
        ws_info.column_dimensions["B"].width = 110
        for row in ws_info.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
    return xlsx_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    out = sys.argv[2] if len(sys.argv) > 2 else None

    def show(done, total, elapsed):
        print(f"  {done}/{total} 完了（{elapsed:.0f}秒経過）", flush=True)

    path, df = run(sys.argv[1], out, on_progress=show)
    print(f"\nExcel出力: {path}")
    print(df["総合判定"].value_counts().to_string())
